import io
import json
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

import pandas as pd
from fastapi import UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.ext.asyncio import AsyncSession
from exceptions.exception import ServiceException, ValidationException, ResourceNotFoundException
from config.error_codes import ErrorCode
from module_admin.entity.do.enums import StatusEnum
from module_admin.entity.do.protocol_do import ProtocolDO
from module_admin.entity.vo.common_vo import CrudResponseModel
from module_admin.entity.vo.protocol_vo import (
    AddProtocolModel,
    ProtocolModel,
    ProtocolPageQueryModel,
    DeleteProtocolModel,
)
from utils.common_util import CamelCaseUtil
from utils.page_util import PageResponseModel
from utils.io_toolkit.models import ImportResult, ImportStatus
from utils.io_toolkit import BatchActionExecutor, BatchActionTask
from utils.transaction_manager import transactional
from ports import IProtocolRepository


PROTOCOL_TYPE_MAPPING = {
    'lan': 'LAN',
    'ethernet': 'LAN',
    'rs422': 'RS422',
    'rs485': 'RS485',
    'can': 'CAN',
    '1553b': '1553B',
}

PROTOCOL_DISPLAY_NAME = {
    'LAN': '以太网',
    'RS422': 'RS422',
    'RS485': 'RS485',
    'CAN': 'CAN',
    '1553B': '1553B',
}

# 不同协议类型对应的“报文配置”区域字段定义，中文名称需与前端协议模板中的标签保持一致
HEADER_FIELDS_BY_PROTOCOL = {
    'LAN': [
        ('sender', '发送方'),
        ('receiver', '接收方'),
        ('frequency', '传输频率'),
        ('speed', '传输速率/bps'),
        ('method', '传输方式'),
        ('port', '端口号'),
        ('frameLength', '帧长度/Byte'),
        ('errorHandling', '错误处理'),
    ],
    'RS422': [
        ('sender', '发送方'),
        ('receiver', '接收方'),
        ('frequency', '传输频率'),
        ('speed', '传输速率/bps'),
        ('method', '传输方式'),
        ('sendDuration', '发送时长/ms'),
        ('frameLength', '帧长度/Byte'),
        ('errorHandling', '错误处理'),
    ],
    'RS485': [
        ('sender', '发送方'),
        ('receiver', '接收方'),
        ('frequency', '传输频率'),
        ('speed', '传输速率/bps'),
        ('method', '传输方式'),
        ('sendDuration', '发送时长/ms'),
        ('frameLength', '帧长度/Byte'),
        ('errorHandling', '错误处理'),
    ],
    'CAN': [
        ('sender', '发送方'),
        ('receiver', '接收方'),
        ('frequency', '传输频率'),
        ('speed', '传输速率/bps'),
        ('method', '传输方式'),
        ('canId', 'ID'),
        ('frameLength', '帧长度/Byte'),
        ('errorHandling', '错误处理'),
    ],
    '1553B': [
        ('sender', '发送方'),
        ('receiver', '接收方'),
        ('frequency', '传输频率'),
        ('speed', '传输速率/bps'),
        ('method', '传输方式'),
        ('subAddress', '子地址'),
        ('frameLength', '帧长度/word'),
        ('errorHandling', '错误处理'),
    ],
}

# 通用头部字段定义，兼容旧模板
HEADER_FIELDS = [
    ('sender', '发送方'),
    ('receiver', '接收方'),
    ('frequency', '传输频率'),
    ('baudRate', '传输速率/bps'),
    ('method', '传输方式'),
    ('duration', '发送时长/ms'),
    ('frameLength', '帧长度/Byte'),
    ('errorHandling', '错误处理'),
]

NON_EDITABLE_HEADER_KEYS = {'sender', 'receiver'}

HEADER_NUMERIC_KEYS = {'baudRate', 'duration', 'frameLength', 'port', 'sendDuration'}
HEADER_SHEET_NAMES = ['报文配置', '报文头', 'header', '基础配置', '协议配置表格']

FIELD_COLUMNS = [
    {'key': 'seq', 'label': '序号', 'width': 10},
    {'key': 'field_name', 'label': '信息名称', 'width': 24},
    {'key': 'byte_count', 'label': '字节数', 'width': 12},
    {'key': 'byte_sequence', 'label': '字节序号', 'width': 14},
    {'key': 'value_range', 'label': '值域及含义', 'width': 32},
    {'key': 'unit', 'label': '量纲', 'width': 12},
    {'key': 'data_type', 'label': '数据类型', 'width': 14},
    {'key': 'scale', 'label': '比例尺', 'width': 12},
    {'key': 'remark', 'label': '备注', 'width': 24},
]

FIELD_SHEET_NAMES = ['字段列表', '字段定义', 'fields', '报文字段', '字段列表表格']

FIELD_COLUMN_ALIASES: Dict[str, List[str]] = {
    'field_code': ['字段标识', '字段ID', '字段编号', 'fieldCode', 'field_id'],
    'field_name': ['信息名称', '字段名称', 'name', 'fieldName'],
    'byte_count': ['字节数', 'byteCount'],
    'byte_sequence': ['字节序号', 'byteSequence'],
    'value_range': ['值域及含义', '值域', 'valueRange', 'meaning'],
    'unit': ['量纲', '单位', 'unit'],
    'data_type': ['数据类型', '类型', 'dataType'],
    'scale': ['比例尺', 'scale'],
    'parent_code': ['父级标识', '父级ID', 'parentId', 'parentCode'],
    'remark': ['备注', 'remark', 'note'],
}

DATA_TYPE_OPTIONS = ['UINT8', 'UINT16', 'UINT32', 'INT8', 'INT16', 'INT32', 'FLOAT', 'DOUBLE', 'STRING']
DATA_TYPE_ALIASES = {
    'uint8': 'UINT8',
    'u8': 'UINT8',
    'byte': 'UINT8',
    '无符号8位': 'UINT8',
    'uint16': 'UINT16',
    'u16': 'UINT16',
    'uint32': 'UINT32',
    'u32': 'UINT32',
    'int8': 'INT8',
    'i8': 'INT8',
    'int16': 'INT16',
    'i16': 'INT16',
    'int32': 'INT32',
    'i32': 'INT32',
    'float': 'FLOAT',
    'f32': 'FLOAT',
    '单精度': 'FLOAT',
    'double': 'DOUBLE',
    'dbl': 'DOUBLE',
    'f64': 'DOUBLE',
    '双精度': 'DOUBLE',
    'string': 'STRING',
    'str': 'STRING',
    '文本': 'STRING',
    '字符串': 'STRING',
}

DEFAULT_HEADER_SAMPLE = {
    'sender': '',
    'receiver': '',
    'frequency': '单次',
    'baudRate': 115200,
    'method': '422',
    'duration': 1000,
    'frameLength': 128,
    'errorHandling': '重传',
}

DEFAULT_HEADER_SAMPLE_EXTRA = {
    'port': '8080',
    'sendDuration': 1000,
    'subAddress': '01',
    'canId': '0x01',
}

DEFAULT_FIELD_SAMPLE = [
    {
        'seq': 1,
        'field_name': '同步码1',
        'byte_count': 1,
        'byte_sequence': '0',
        'value_range': '0xEB',
        'unit': '',
        'data_type': 'UINT8',
        'scale': '1',
        'remark': '',
    },
    {
        'seq': 2,
        'field_name': '同步码2',
        'byte_count': 1,
        'byte_sequence': '1',
        'value_range': '0x90',
        'unit': '',
        'data_type': 'UINT8',
        'scale': '1',
        'remark': '',
    },
    {
        'seq': 3,
        'field_name': '报文字节数',
        'byte_count': 2,
        'byte_sequence': '2~3',
        'value_range': '0x0100',
        'unit': '',
        'data_type': 'UINT16',
        'scale': '1',
        'remark': '',
    },
]

CONFIG_TABLE_LAYOUT = {
    'LAN': [
        [('sender', '发送方')],
        [('receiver', '接收方')],
        [('frequency', '传输频率'), ('speed', '传输速率/bps')],
        [('method', '传输方式'), ('port', '端口号')],
        [('frameLength', '帧长度/Byte'), ('errorHandling', '错误处理')],
    ],
    'RS422': [
        [('sender', '发送方')],
        [('receiver', '接收方')],
        [('frequency', '传输频率'), ('speed', '传输速率/bps')],
        [('method', '传输方式'), ('sendDuration', '发送时长/ms')],
        [('frameLength', '帧长度/Byte'), ('errorHandling', '错误处理')],
    ],
    'RS485': [
        [('sender', '发送方')],
        [('receiver', '接收方')],
        [('frequency', '传输频率'), ('speed', '传输速率/bps')],
        [('method', '传输方式'), ('sendDuration', '发送时长/ms')],
        [('frameLength', '帧长度/Byte'), ('errorHandling', '错误处理')],
    ],
    'CAN': [
        [('sender', '发送方')],
        [('receiver', '接收方')],
        [('frequency', '传输频率'), ('speed', '传输速率/bps')],
        [('method', '传输方式'), ('canId', 'ID')],
        [('frameLength', '帧长度/Byte'), ('errorHandling', '错误处理')],
    ],
    '1553B': [
        [('sender', '发送方')],
        [('receiver', '接收方')],
        [('frequency', '传输频率'), ('speed', '传输速率/bps')],
        [('method', '传输方式'), ('subAddress', '子地址')],
        [('frameLength', '帧长度/word'), ('errorHandling', '错误处理')],
    ],
}

HEADER_LABEL_TO_KEY = {
    '协议类型': 'protocolType',
    'ProtocolType': 'protocolType',
    'Protocol Type': 'protocolType',
    '发送方': 'sender',
    '接收方': 'receiver',
    '传输频率': 'frequency',
    '发送频率': 'frequency',
    '传输速率/bps': 'speed',
    '波特率': 'baudRate',
    '传输方式': 'method',
    '发送方式': 'method',
    '端口号': 'port',
    '发送时长/ms': 'sendDuration',
    '发送时长': 'sendDuration',
    'Duration': 'duration',
    '帧长度/Byte': 'frameLength',
    '帧长度/byte': 'frameLength',
    '帧长度/word': 'frameLength',
    '帧长度/Word': 'frameLength',
    '帧长': 'frameLength',
    '子地址': 'subAddress',
    'ID': 'canId',
    'Id': 'canId',
    '错误处理': 'errorHandling',
}

ALL_HEADER_KEYS = (
    {key for key, _ in HEADER_FIELDS}
    | {key for fields in HEADER_FIELDS_BY_PROTOCOL.values() for key, _ in fields}
    | {'protocolType'}
)


class ProtocolService:
    """
    协议管理模块服务层
    """

    def __init__(self, protocol_repo: IProtocolRepository, query_db: Optional[AsyncSession] = None):
        self.protocol_repo = protocol_repo
        self.query_db = query_db

    async def get_protocol_list(self, query_object: ProtocolPageQueryModel, is_page: bool = False):
        return await self.get_protocol_list_services(query_object, is_page=is_page)

    async def add_protocol(self, page_object: ProtocolModel, operator: Optional[str] = None):
        return await self.add_protocol_services(page_object, operator=operator)

    async def edit_protocol(self, page_object: ProtocolModel, operator: Optional[str] = None):
        return await self.edit_protocol_services(page_object, operator=operator)

    async def delete_protocol(self, page_object: DeleteProtocolModel, operator: Optional[str] = None):
        return await self.delete_protocol_services(page_object, operator=operator)

    async def delete_protocol_batch(self, protocol_ids: List[int], operator: Optional[str] = None):
        return await self.delete_protocol_batch_services(protocol_ids=protocol_ids, operator=operator or '')

    async def get_protocol_detail(self, protocol_id: int):
        return await self.protocol_detail_services(protocol_id)

    @staticmethod
    def __normalize_protocol_config_for_frontend(raw):
        """
        将协议配置规范化为前端编辑页可直接渲染的结构：
        - 头部字段平铺在根对象
        - 字段列表放在 fields
        兼容历史/导入时的 headerConfig / fieldsConfig / messageConfig 等嵌套结构。
        """
        if not isinstance(raw, dict):
            return raw if raw else {}

        # 已经是前端结构
        if isinstance(raw.get('fields'), list):
            return raw

        header = raw.get('headerConfig') or raw.get('header') or {}
        if not isinstance(header, dict):
            header = {}

        message_cfg = raw.get('messageConfig') if isinstance(raw.get('messageConfig'), dict) else {}
        if isinstance(message_cfg.get('header'), dict) and not header:
            header = message_cfg.get('header')

        fields = raw.get('fieldsConfig')
        if not isinstance(fields, list):
            fields = message_cfg.get('fields')
        if not isinstance(fields, list):
            fields = []

        # 如果没有解析到 header/fields，则直接返回原始对象，避免误覆盖
        if not header and not fields:
            return raw

        merged = {**header, 'fields': fields}
        return merged

    async def get_protocol_list_services(self, query_object: ProtocolPageQueryModel, is_page: bool = False):
        """
        获取协议列表信息service

        :param query_object: 查询参数对象
        :param is_page: 是否开启分页
        :return: 协议列表信息对象
        """
        if is_page:
            protocol_list_result = await self.protocol_repo.paginate_with_filters(
                page=query_object.page,
                page_size=query_object.page_size,
                protocol_name=query_object.protocol_name,
                protocol_type=query_object.protocol_type,
                begin_time=query_object.begin_time,
                end_time=query_object.end_time,
            )
            items = CamelCaseUtil.transform_result(protocol_list_result.items)
            for row in items:
                if isinstance(row, dict) and 'protocolConfig' in row:
                    row['protocolConfig'] = self.__normalize_protocol_config_for_frontend(row.get('protocolConfig'))
            return PageResponseModel(**{**protocol_list_result.model_dump(), 'items': items})
        else:
            protocol_list_result = await self.protocol_repo.find_all_active()
            return CamelCaseUtil.transform_result(protocol_list_result)

    @transactional()
    async def add_protocol_services(
        self,
        page_object: ProtocolModel,
        operator: Optional[str] = None,
    ):
        """
        新增协议信息service

        :param page_object: 新增协议对象
        :return: 新增协议校验结果
        """
        now = datetime.now()
        effective_operator = operator or page_object.create_by or page_object.update_by or 'system'
        if not page_object.create_by:
            page_object.create_by = effective_operator
        if not page_object.update_by:
            page_object.update_by = effective_operator
        if not page_object.create_time:
            page_object.create_time = now
        if not page_object.update_time:
            page_object.update_time = now

        db_protocol = ProtocolDO(**page_object.model_dump())
        add_result = await self.protocol_repo.add(db_protocol)
        return CrudResponseModel(is_success=True, message='新增成功', result={'protocolId': add_result.protocol_id})

    @transactional()
    async def edit_protocol_services(
        self,
        page_object: ProtocolModel,
        operator: Optional[str] = None,
    ):
        """
        编辑协议信息service

        :param page_object: 编辑协议对象
        :return: 编辑协议校验结果
        """
        protocol = await self.protocol_repo.find_by_id(page_object.protocol_id)
        if protocol:
            if operator and not page_object.update_by:
                page_object.update_by = operator
            if not page_object.update_time:
                page_object.update_time = datetime.now()
            update_data = page_object.model_dump(exclude_unset=True)
            for key, value in update_data.items():
                if hasattr(protocol, key) and value is not None:
                    setattr(protocol, key, value)
            await self.protocol_repo.update(protocol)
        return CrudResponseModel(is_success=True, message='更新成功', result={'protocolId': protocol.protocol_id})

    @transactional()
    async def delete_protocol_services(
        self,
        page_object: DeleteProtocolModel,
        operator: Optional[str] = None,
    ):
        """
        删除协议信息service

        :param page_object: 删除协议对象
        :return: 删除协议校验结果
        """
        if operator and not page_object.update_by:
            page_object.update_by = operator
        protocol_id_list = [int(id) for id in page_object.protocol_ids.split(',') if id]
        await self.protocol_repo.soft_delete_batch(protocol_id_list, page_object.update_by or '')
        return CrudResponseModel(is_success=True, message='删除成功')

    async def delete_protocol_batch_services(
        self,
        protocol_ids: List[int],
        operator: str,
    ) -> Dict[str, Any]:
        """
        批量删除协议，使用 BatchActionExecutor 统一结果结构
        """
        if not protocol_ids:
            raise ValidationException(
                error_code=ErrorCode.INVALID_INPUT, field='protocol_ids', user_message='请选择要删除的协议'
            )

        normalized_ids: List[int] = []
        invalid_ids: List[str] = []
        for pid in protocol_ids:
            try:
                normalized_ids.append(int(pid))
            except Exception:
                invalid_ids.append(str(pid))

        executor = BatchActionExecutor(concurrency=1, enable_dedup=True)
        async def _delete_one(pid: int):
            await self.protocol_repo.soft_delete_batch([pid], operator)
            await self.protocol_repo.commit()

        tasks = [
            BatchActionTask(
                id=pid,
                handler=_delete_one,
                args=(pid,),
                timeout=5,
                max_retries=0,
                dedup_key=str(pid),
            )
            for pid in normalized_ids
        ]

        summary = await executor.run(tasks)
        result = summary.as_dict()
        if invalid_ids:
            # 记录无法解析的ID
            result.setdefault('failed', []).extend(
                [
                    {'id': bad_id, 'success': False, 'data': None, 'error': '非法ID', 'attempts': 0}
                    for bad_id in invalid_ids
                ]
            )
        return result

    async def protocol_detail_services(self, protocol_id: int):
        """
        获取协议详细信息service

        :param protocol_id: 协议ID
        :return: 协议ID对应的信息
        """
        protocol = await self.protocol_repo.find_by_id(protocol_id)
        if protocol:
            result = CamelCaseUtil.transform_result(protocol)
            if isinstance(result, dict) and 'protocolConfig' in result:
                result['protocolConfig'] = self.__normalize_protocol_config_for_frontend(result.get('protocolConfig'))
        else:
            raise ResourceNotFoundException(
                user_message='协议不存在',
                error_code=ErrorCode.PROTOCOL_NOT_FOUND,
                developer_message=f'protocol_id={protocol_id} not found',
            )

        return result

    async def export_protocol_config_services(self, protocol_id: int):
        """
        导出协议配置service

        :param protocol_id: 协议ID
        :return: 包含协议名称和Excel内容的字典
        """
        protocol = await self.protocol_repo.find_by_id(protocol_id)
        if not protocol:
            raise ResourceNotFoundException(error_code=ErrorCode.RESOURCE_NOT_FOUND, user_message='协议不存在')

        protocol_data = CamelCaseUtil.transform_result(protocol)
        protocol_name = protocol_data.get('protocolName', '协议配置')

        # 兼容 protocol_config 存为 JSON 字符串的情况
        def ensure_dict(obj):
            if isinstance(obj, dict):
                return obj
            if isinstance(obj, str) and obj.strip():
                try:
                    import json

                    loaded = json.loads(obj)
                    return loaded if isinstance(loaded, dict) else {}
                except Exception:
                    return {}
            return {}

        protocol_config = ensure_dict(protocol_data.get('protocolConfig') or {})
        nested_protocol_config = ensure_dict(protocol_config.get('protocolConfig'))
        protocol_type = (
            protocol_data.get('protocolType')
            or (protocol_config.get('protocolType') if isinstance(protocol_config, dict) else '')
            or (nested_protocol_config.get('protocolType') if isinstance(nested_protocol_config, dict) else '')
        )

        def resolve_header_config():
            # 兼容不同存储结构：
            # 1) 顶层 headerConfig
            # 2) protocolConfig.headerConfig
            # 3) protocolConfig.header (部分保存/导入逻辑使用 header 字段)
            # 4) protocolConfig.protocolConfig 内层存储
            # 5) 平铺在 protocolConfig（除去字段列表）
            # 6) messageConfig 结构（历史设备协议或导入数据）
            def merge_header_sources(sources):
                merged = {}
                for source in sources:
                    if not isinstance(source, dict):
                        continue
                    for key, val in source.items():
                        # 跳过明显不是头部配置的字段
                        if key in {'fields', 'fieldsConfig', 'protocolConfig', 'header', 'headerConfig'}:
                            continue
                        # 避免用空值覆盖已有值
                        if val in (None, '') and merged.get(key) not in (None, ''):
                            continue
                        merged[key] = val
                return merged

            header_sources = [
                ensure_dict(protocol_data.get('headerConfig')),
                ensure_dict(protocol_data.get('header')),
            ]

            message_config = ensure_dict(protocol_data.get('messageConfig'))
            if isinstance(message_config, dict):
                header_sources.append(ensure_dict(message_config.get('header')))
                header_sources.append(
                    {
                        k: v
                        for k, v in message_config.items()
                        if k not in {'fields', 'fieldsConfig', 'protocolConfig', 'header', 'headerConfig'}
                    }
                )

            if isinstance(protocol_config, dict):
                header_sources.append(ensure_dict(protocol_config.get('headerConfig')))
                header_sources.append(ensure_dict(protocol_config.get('header')))
                header_sources.append(ensure_dict(protocol_config.get('messageConfig', {}).get('header')))
                # 将平铺在 protocolConfig 的字段纳入
                header_sources.append(
                    {
                        k: v
                        for k, v in protocol_config.items()
                        if k not in {'fields', 'fieldsConfig', 'protocolConfig', 'header', 'headerConfig'}
                    }
                )

            if isinstance(nested_protocol_config, dict):
                header_sources.append(ensure_dict(nested_protocol_config.get('headerConfig')))
                header_sources.append(ensure_dict(nested_protocol_config.get('header')))
                header_sources.append(ensure_dict(nested_protocol_config.get('messageConfig', {}).get('header')))
                header_sources.append(
                    {
                        k: v
                        for k, v in nested_protocol_config.items()
                        if k not in {'fields', 'fieldsConfig', 'protocolConfig', 'header', 'headerConfig'}
                    }
                )

            header_cfg = merge_header_sources(header_sources)
            return header_cfg or {}

        def resolve_fields_config():
            # 兼容不同存储结构：顶层 fieldsConfig、protocolConfig.fieldsConfig、protocolConfig.fields
            message_config = ensure_dict(protocol_data.get('messageConfig'))
            protocol_message_config = (
                ensure_dict(protocol_config.get('messageConfig')) if isinstance(protocol_config, dict) else {}
            )
            nested_message_config = (
                ensure_dict(nested_protocol_config.get('messageConfig'))
                if isinstance(nested_protocol_config, dict)
                else {}
            )
            candidates = [
                protocol_data.get('fieldsConfig'),
                protocol_data.get('fields'),
                message_config.get('fields'),
                protocol_config.get('fieldsConfig') if isinstance(protocol_config, dict) else None,
                protocol_config.get('fields') if isinstance(protocol_config, dict) else None,
                protocol_message_config.get('fields'),
                nested_protocol_config.get('fieldsConfig') if isinstance(nested_protocol_config, dict) else None,
                nested_protocol_config.get('fields') if isinstance(nested_protocol_config, dict) else None,
                nested_message_config.get('fields'),
            ]
            for item in candidates:
                if item:
                    return item
            # 确保至少返回空列表
            return []

        header_config = resolve_header_config()
        fields_config = resolve_fields_config()

        def flatten_configs(configs):
            """
            将多个配置源拍平，便于取值，跳过字段列表等大块内容。
            保留第一个非空值，避免被空值覆盖。
            """
            merged: Dict[str, Any] = {}
            visited: set[int] = set()

            def walk(obj: Dict[str, Any]):
                obj_id = id(obj)
                if obj_id in visited:
                    return
                visited.add(obj_id)
                for k, v in obj.items():
                    if isinstance(v, dict):
                        walk(v)
                    if k in {'fields', 'fieldsConfig', 'protocolConfig', 'header', 'headerConfig'}:
                        continue
                    if k not in merged or merged[k] in (None, ''):
                        merged[k] = v

            for cfg in configs:
                if isinstance(cfg, dict):
                    walk(cfg)
            return merged

        flattened_header_source = flatten_configs(
            [
                protocol_data,
                header_config,
                protocol_config,
                nested_protocol_config,
            ]
        )

        # 标准化协议类型
        normalized_protocol_type = PROTOCOL_TYPE_MAPPING.get(protocol_type.lower() if protocol_type else '', None)
        header_fields = HEADER_FIELDS_BY_PROTOCOL.get(normalized_protocol_type, HEADER_FIELDS)
        display_name = PROTOCOL_DISPLAY_NAME.get(normalized_protocol_type, normalized_protocol_type or '通用')
        protocol_type_value = normalized_protocol_type or (protocol_type or '通用')
        protocol_type_display = display_name or protocol_type_value or '通用'
        layout_rows = CONFIG_TABLE_LAYOUT.get(normalized_protocol_type, self.__build_default_layout(header_fields))

        workbook = Workbook()
        header_sheet = workbook.active
        header_sheet.title = '协议配置表格'

        label_fill = PatternFill(start_color='f7f8fa', end_color='f7f8fa', fill_type='solid')
        value_fill = PatternFill(start_color='ffffff', end_color='ffffff', fill_type='solid')
        head_font = Font(bold=True)
        thin_border = Border(
            left=Side(border_style='thin', color='d9d9d9'),
            right=Side(border_style='thin', color='d9d9d9'),
            top=Side(border_style='thin', color='d9d9d9'),
            bottom=Side(border_style='thin', color='d9d9d9'),
        )

        # 写入协议类型元数据行，便于导出文件可直接被导入解析
        header_sheet.append(['协议类型', '协议显示名', '', '', ''])
        header_sheet.append([protocol_type_value, protocol_type_display, '', '', ''])
        # 空行分隔视觉区域
        header_sheet.append(['', '', '', '', ''])
        # 隐藏协议元数据行，避免导出时在视觉上多出额外的提示行
        header_sheet.row_dimensions[1].hidden = True
        header_sheet.row_dimensions[2].hidden = True
        header_sheet.row_dimensions[3].hidden = True

        header_sheet.freeze_panes = f'A{header_sheet.max_row + 1}'

        first_data_row = header_sheet.max_row + 1
        last_row = first_data_row + len(layout_rows) - 1
        header_sheet.merge_cells(start_row=first_data_row, start_column=1, end_row=last_row, end_column=1)
        protocol_cell = header_sheet.cell(row=first_data_row, column=1, value=display_name)
        protocol_cell.alignment = Alignment(horizontal='center', vertical='center')
        protocol_cell.fill = label_fill
        protocol_cell.font = head_font

        def apply_range_border(sheet, r1, r2, c1, c2):
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    sheet.cell(row=r, column=c).border = thin_border

        header_key_alias = {
            # 优先展示带 Label 的中文值，再回退到原始值/别名
            # 前端使用 frequencyLabel 保存中文标签，frequency 保存值
            'frequency': ['frequencyLabel', 'frequency_label', 'frequency'],
            # 前端使用 speedLabel 保存中文标签，speed 保存值；后端可能用 baudRate
            'speed': [
                'speedLabel',
                'speed_label',
                'speed',
                'baudRate',
                'baud_rate',
                'baudRateLabel',
                'baud_rate_label',
            ],
            # 前端使用 methodLabel 保存中文标签，method 保存值
            'method': ['methodLabel', 'method_label', 'method'],
            # 前端使用 errorHandle/errorHandleLabel，后端导出模板用 errorHandling
            'errorHandling': [
                'errorHandleLabel',
                'error_handle_label',
                'errorHandle',
                'error_handle',
                'errorHandling',
                'errorHandlingLabel',
                'error_handling_label',
            ],
            # 前端使用 sendDuration，后端可能用 duration
            'sendDuration': ['sendDuration', 'send_duration', 'duration'],
            # 帧长度
            'frameLength': ['frameLength', 'frameLengthLabel', 'frame_length', 'frame_length_label'],
            # 端口号
            'port': ['port', 'portNo', 'portNumber', 'port_no', 'port_number'],
        }

        def pick_value(src: Dict, candidates: List[str]):
            for name in candidates:
                if name not in src:
                    continue
                val = src.get(name, '')
                if isinstance(val, dict):
                    # 兼容存储为 {label, value} 或 {text, value}
                    extracted = val.get('label') or val.get('text') or val.get('value') or ''
                    if extracted not in (None, ''):
                        return extracted
                    continue
                # 跳过空值，继续查找下一个候选项
                if val in (None, ''):
                    continue
                return val
            return ''

        def get_header_value(key):
            """从header_config/protocol_config中获取值"""
            if key == 'protocolType':
                return protocol_type or protocol_config.get('protocolType', '')

            candidates = header_key_alias.get(key, []) + [key, f'{key}Label']

            if isinstance(header_config, dict):
                val = pick_value(header_config, candidates)
                if val not in (None, ''):
                    return val

            if isinstance(protocol_config, dict):
                val = pick_value(protocol_config, candidates)
                if val not in (None, ''):
                    return val

            # 深度兼容：在拍平后的配置源中兜底取值
            val = pick_value(flattened_header_source, candidates)
            if val not in (None, ''):
                return val

            return ''

        for row_offset, pair in enumerate(layout_rows):
            current_row = first_data_row + row_offset
            if pair:
                key1, label1 = pair[0]
                c_label1 = header_sheet.cell(row=current_row, column=2, value=label1)
                c_label1.alignment = Alignment(horizontal='center', vertical='center')
                c_label1.fill = label_fill
                c_label1.font = head_font

                c_val1 = header_sheet.cell(row=current_row, column=3, value=get_header_value(key1))
                c_val1.fill = value_fill

            if len(pair) > 1:
                key2, label2 = pair[1]
                c_label2 = header_sheet.cell(row=current_row, column=4, value=label2)
                c_label2.alignment = Alignment(horizontal='center', vertical='center')
                c_label2.fill = label_fill
                c_label2.font = head_font
                c_val2 = header_sheet.cell(row=current_row, column=5, value=get_header_value(key2))
                c_val2.fill = value_fill

            # 如果当前行只有一个字段，横向合并
            if len(pair) == 1 or (len(pair) > 1 and not pair[1][0]):
                header_sheet.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=5)

        header_sheet.column_dimensions['A'].width = 12
        header_sheet.column_dimensions['B'].width = 18
        header_sheet.column_dimensions['C'].width = 24
        header_sheet.column_dimensions['D'].width = 18
        header_sheet.column_dimensions['E'].width = 24
        apply_range_border(header_sheet, first_data_row, last_row, 1, 5)

        # 绘制字段列表表格（主 sheet，保持原有布局）
        field_table_start = last_row + 2
        for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
            cell = header_sheet.cell(row=field_table_start, column=column_index, value=column_meta['label'])
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = label_fill
            cell.font = head_font
            header_sheet.column_dimensions[get_column_letter(column_index)].width = column_meta['width']

        field_key_mapping = {
            'seq': 'seq',
            'field_name': 'fieldName',
            'byte_count': 'byteCount',
            'byte_sequence': 'byteSequence',
            'value_range': 'valueRange',
            'unit': 'unit',
            'data_type': 'dataType',
            'scale': 'scale',
            'remark': 'remark',
        }

        def resolve_seq(field_data: dict, default_seq: int):
            if isinstance(field_data, dict):
                seq_val = field_data.get('seq')
                if seq_val not in (None, ''):
                    return seq_val
            return default_seq

        if fields_config and isinstance(fields_config, list):
            for row_index, field_data in enumerate(fields_config, start=field_table_start + 1):
                for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
                    key = column_meta['key']
                    mapped_key = field_key_mapping.get(key, key)
                    if key == 'seq':
                        value = resolve_seq(field_data, row_index - field_table_start)
                    else:
                        value = field_data.get(mapped_key, '') if isinstance(field_data, dict) else ''
                    data_cell = header_sheet.cell(row=row_index, column=column_index, value=value)
                    data_cell.fill = value_fill
            apply_range_border(
                header_sheet, field_table_start, field_table_start + len(fields_config), 1, len(FIELD_COLUMNS)
            )
        else:
            apply_range_border(header_sheet, field_table_start, field_table_start, 1, len(FIELD_COLUMNS))

        # 额外生成“字段列表表格”sheet，保持与模板一致，确保导出文件可直接导入
        field_sheet = workbook.create_sheet(title='字段列表表格')
        field_sheet.freeze_panes = 'A2'
        header_fill = PatternFill(start_color='f2f2f2', end_color='f2f2f2', fill_type='solid')
        for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
            cell = field_sheet.cell(row=1, column=column_index, value=column_meta['label'])
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = header_fill
            field_sheet.column_dimensions[get_column_letter(column_index)].width = column_meta['width']

        target_fields = fields_config if fields_config and isinstance(fields_config, list) else []
        for row_index, field_data in enumerate(target_fields, start=2):
            for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
                key = column_meta['key']
                mapped_key = field_key_mapping.get(key, key)
                if key == 'seq':
                    value = resolve_seq(field_data, row_index - 1)
                else:
                    value = field_data.get(mapped_key, '') if isinstance(field_data, dict) else ''
                field_sheet.cell(row=row_index, column=column_index, value=value)

        data_type_column_index = next(
            (index for index, column_meta in enumerate(FIELD_COLUMNS, start=1) if column_meta['key'] == 'data_type'),
            None,
        )
        if data_type_column_index:
            data_type_column_letter = get_column_letter(data_type_column_index)

            dv_header = DataValidation(
                type='list',
                formula1=f'"{",".join(DATA_TYPE_OPTIONS)}"',
                allow_blank=True,
                showErrorMessage=True,
            )
            header_start_row = field_table_start + 1
            dv_header.add(f'{data_type_column_letter}{header_start_row}:{data_type_column_letter}1048576')
            header_sheet.add_data_validation(dv_header)

            dv_fields = DataValidation(
                type='list',
                formula1=f'"{",".join(DATA_TYPE_OPTIONS)}"',
                allow_blank=True,
                showErrorMessage=True,
            )
            dv_fields.add(f'{data_type_column_letter}2:{data_type_column_letter}1048576')
            field_sheet.add_data_validation(dv_fields)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return {'protocol_name': protocol_name, 'content': buffer.getvalue()}

    async def get_protocol_import_template_services(
        self, protocol_type: Optional[str] = None, protocolType: Optional[str] = None, **kwargs
    ):
        """
        获取协议导入模板excel内容
        :param protocol_type: 协议类型 (lan/rs422/rs485/can/1553b)
        """
        # 兼容不同的参数命名
        protocol_type = protocol_type or protocolType

        normalized_protocol_type = PROTOCOL_TYPE_MAPPING.get(protocol_type.lower() if protocol_type else '', None)
        header_fields = HEADER_FIELDS_BY_PROTOCOL.get(normalized_protocol_type, HEADER_FIELDS)
        display_name = PROTOCOL_DISPLAY_NAME.get(normalized_protocol_type, normalized_protocol_type or '通用')
        layout_rows = CONFIG_TABLE_LAYOUT.get(normalized_protocol_type, self.__build_default_layout(header_fields))

        workbook = Workbook()
        header_sheet = workbook.active
        header_sheet.title = '协议配置表格'

        label_fill = PatternFill(start_color='f7f8fa', end_color='f7f8fa', fill_type='solid')
        value_fill = PatternFill(start_color='ffffff', end_color='ffffff', fill_type='solid')
        head_font = Font(bold=True)
        thin_border = Border(
            left=Side(border_style='thin', color='d9d9d9'),
            right=Side(border_style='thin', color='d9d9d9'),
            top=Side(border_style='thin', color='d9d9d9'),
            bottom=Side(border_style='thin', color='d9d9d9'),
        )

        # 空表头行，保持与前端所见格子一致（无标题）
        header_sheet.append(['', '', '', '', ''])
        header_sheet.freeze_panes = 'A2'

        first_data_row = 2
        last_row = first_data_row + len(layout_rows) - 1
        header_sheet.merge_cells(start_row=first_data_row, start_column=1, end_row=last_row, end_column=1)
        protocol_cell = header_sheet.cell(row=first_data_row, column=1, value=display_name)
        protocol_cell.alignment = Alignment(horizontal='center', vertical='center')
        protocol_cell.fill = label_fill
        protocol_cell.font = head_font

        def apply_range_border(sheet, r1, r2, c1, c2):
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    sheet.cell(row=r, column=c).border = thin_border

        def configure_header_value_cell(cell, key):
            sample_value = '' if key in NON_EDITABLE_HEADER_KEYS else self.__header_sample_value(key)
            cell.value = sample_value
            cell.fill = value_fill
            cell.protection = Protection(locked=key in NON_EDITABLE_HEADER_KEYS)

        for row_offset, pair in enumerate(layout_rows):
            current_row = first_data_row + row_offset
            if pair:
                key1, label1 = pair[0]
                c_label1 = header_sheet.cell(row=current_row, column=2, value=label1)
                c_label1.alignment = Alignment(horizontal='center', vertical='center')
                c_label1.fill = label_fill
                c_label1.font = head_font

                c_val1 = header_sheet.cell(row=current_row, column=3)
                configure_header_value_cell(c_val1, key1)
            if len(pair) > 1:
                key2, label2 = pair[1]
                c_label2 = header_sheet.cell(row=current_row, column=4, value=label2)
                c_label2.alignment = Alignment(horizontal='center', vertical='center')
                c_label2.fill = label_fill
                c_label2.font = head_font
                c_val2 = header_sheet.cell(row=current_row, column=5)
                configure_header_value_cell(c_val2, key2)

            # 如果当前行只有 label1 / value1，也做横向合并
            if len(pair) == 1 or (len(pair) > 1 and not pair[1][0]):
                header_sheet.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=5)

        header_sheet.column_dimensions['A'].width = 12
        header_sheet.column_dimensions['B'].width = 18
        header_sheet.column_dimensions['C'].width = 24
        header_sheet.column_dimensions['D'].width = 18
        header_sheet.column_dimensions['E'].width = 24
        apply_range_border(header_sheet, first_data_row, last_row, 1, 5)

        # 在同一张表下继续绘制字段列表表格，保持与示例一致
        field_table_start = last_row + 2
        for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
            cell = header_sheet.cell(row=field_table_start, column=column_index, value=column_meta['label'])
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = label_fill
            cell.font = head_font
            header_sheet.column_dimensions[get_column_letter(column_index)].width = column_meta['width']

        for row_index, row_data in enumerate(DEFAULT_FIELD_SAMPLE, start=field_table_start + 1):
            for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
                value = self.__field_value(column_meta['key'], row_data, row_index - field_table_start)
                data_cell = header_sheet.cell(row=row_index, column=column_index, value=value)
                data_cell.fill = value_fill
                data_cell.protection = Protection(locked=False)
        apply_range_border(
            header_sheet, field_table_start, field_table_start + len(DEFAULT_FIELD_SAMPLE), 1, len(FIELD_COLUMNS)
        )
        header_sheet.protection.sheet = True

        # 同时保留单独的字段列表表，用于导入解析兼容
        field_sheet = workbook.create_sheet(title='字段列表表格')
        field_sheet.freeze_panes = 'A2'
        header_fill = PatternFill(start_color='f2f2f2', end_color='f2f2f2', fill_type='solid')
        for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
            cell = field_sheet.cell(row=1, column=column_index, value=column_meta['label'])
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = header_fill
            field_sheet.column_dimensions[get_column_letter(column_index)].width = column_meta['width']

        for row_index, row_data in enumerate(DEFAULT_FIELD_SAMPLE, start=2):
            for column_index, column_meta in enumerate(FIELD_COLUMNS, start=1):
                value = self.__field_value(column_meta['key'], row_data, row_index - 1)
                field_sheet.cell(row=row_index, column=column_index, value=value)

        data_type_column_index = next(
            (index for index, column_meta in enumerate(FIELD_COLUMNS, start=1) if column_meta['key'] == 'data_type'),
            None,
        )
        if data_type_column_index:
            data_type_column_letter = get_column_letter(data_type_column_index)

            dv_header = DataValidation(
                type='list',
                formula1=f'"{",".join(DATA_TYPE_OPTIONS)}"',
                allow_blank=True,
                showErrorMessage=True,
            )
            header_start_row = field_table_start + 1
            dv_header.add(f'{data_type_column_letter}{header_start_row}:{data_type_column_letter}1048576')
            header_sheet.add_data_validation(dv_header)

            dv_fields = DataValidation(
                type='list',
                formula1=f'"{",".join(DATA_TYPE_OPTIONS)}"',
                allow_blank=True,
                showErrorMessage=True,
            )
            dv_fields.add(f'{data_type_column_letter}2:{data_type_column_letter}1048576')
            field_sheet.add_data_validation(dv_fields)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    async def protocol_import_preview_services(self, file: UploadFile):
        """
        解析导入的协议配置文件
        支持 Excel (.xlsx, .xls, .csv) 和 JSON (.json) 格式
        """
        if not file or not file.filename:
            raise ValidationException(error_code=ErrorCode.INVALID_INPUT, user_message='请选择要导入的文件')
        suffix = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
        if suffix not in ('xlsx', 'xls', 'csv', 'json'):
            raise ValidationException(
                error_code=ErrorCode.INVALID_INPUT, field='file', user_message='仅支持 .xlsx/.xls/.csv/.json 格式的文件'
            )

        content = await file.read()
        if not content:
            raise ValidationException(error_code=ErrorCode.INVALID_INPUT, user_message='文件内容为空')

        # JSON 格式处理
        if suffix == 'json':
            parsed = self.__parse_json_protocol(content)
            parsed['header'] = self.__normalize_header_enums(parsed.get('header') or {}, parsed.get('protocolType'))
            return parsed

        # Excel 格式处理
        sheet_map = self.__read_workbook(content, suffix)
        protocol_type = self.__extract_protocol_type(sheet_map)
        header = self.__normalize_header_enums(self.__extract_header(sheet_map), protocol_type)
        fields = self.__extract_fields(sheet_map)
        if not fields:
            raise ValidationException(
                error_code=ErrorCode.INVALID_INPUT, user_message='文件中未解析到任何字段，请检查模板内容'
            )
        return {'header': header, 'fields': fields, 'protocolType': protocol_type}

    @transactional()
    async def import_protocol_services(self, file: UploadFile, protocol_name: Optional[str], operator: str):
        """
        导入协议配置service
        解析导出的Excel或JSON文件并创建新协议

        :param file: 上传的Excel或JSON文件
        :param protocol_name: 协议名称（可选，若不提供则从文件名提取）
        :param operator: 操作人
        :return: 导入结果
        """
        if not file or not file.filename:
            raise ServiceException(user_message='请选择要导入的文件')

        suffix = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
        if suffix not in ('xlsx', 'xls', 'csv', 'json'):
            raise ServiceException(user_message='仅支持 .xlsx/.xls/.csv/.json 格式的文件')

        content = await file.read()
        if not content:
            raise ServiceException(message='文件内容为空')

        # JSON 格式处理
        if suffix == 'json':
            parsed_data = self.__parse_json_protocol(content)
            protocol_type = parsed_data.get('protocolType', 'LAN')
            header = self.__normalize_header_enums(parsed_data.get('header', {}), protocol_type)
            fields = parsed_data.get('fields', [])
        else:
            # Excel 格式处理
            sheet_map = self.__read_workbook(content, suffix)
            protocol_type = self.__extract_protocol_type(sheet_map)
            header = self.__normalize_header_enums(self.__extract_header(sheet_map), protocol_type)
            fields = self.__extract_fields(sheet_map)

        if not fields:
            raise ServiceException(message='文件中未解析到任何字段，请检查模板内容')

        # 确定协议名称
        if not protocol_name:
            # 从文件名提取，移除后缀和"_配置"后缀
            base_name = file.filename.rsplit('.', 1)[0] if '.' in file.filename else file.filename
            if base_name.endswith('_配置'):
                base_name = base_name[:-3]
            protocol_name = base_name

        # 检查协议名称是否已存在，如果存在则添加时间戳
        existing = await self.protocol_repo.find_by_name(protocol_name)
        if existing:
            import time

            timestamp = int(time.time())
            protocol_name = f'{protocol_name}_{timestamp}'

        # 标准化协议类型
        normalized_type = PROTOCOL_TYPE_MAPPING.get(
            protocol_type.lower() if protocol_type else '', protocol_type or 'LAN'
        )

        # 构建协议配置（与前端编辑页期望结构对齐：头部字段平铺 + fields 列表）
        header_payload = dict(header or {})
        header_payload['protocolType'] = normalized_type
        protocol_config = {
            **header_payload,
            'fields': fields,
            # 兼容导出/旧逻辑：同时保留嵌套结构
            'headerConfig': header_payload,
            'fieldsConfig': fields,
            'messageConfig': {'header': header_payload, 'fields': fields},
        }

        # 创建协议模型
        protocol_model = ProtocolModel(
            protocolName=protocol_name,
            protocolType=normalized_type,
            protocolConfig=protocol_config,
            status=StatusEnum.NORMAL.value,
            createBy=operator,
            updateBy=operator,
        )

        # 保存协议
        db_protocol = ProtocolDO(**protocol_model.model_dump())
        created_protocol = await self.protocol_repo.add(db_protocol)

        return CrudResponseModel(
            is_success=True,
            message='导入成功',
            result={
                'protocolId': created_protocol.protocol_id,
                'protocolName': protocol_name,
                'protocolType': normalized_type,
            },
        )

    async def import_protocol_batch_services(self, files: List[UploadFile], operator: str) -> ImportResult:
        """
        批量导入协议配置

        :param files: 上传的文件列表
        :param operator: 操作人
        :return: ImportResult
        """
        if not files:
            raise ServiceException(message='请选择要导入的文件')

        success_items: List[Dict[str, Any]] = []
        failed_items: List[Dict[str, Any]] = []

        for file in files:
            filename = getattr(file, 'filename', '') or '未知文件'
            try:
                single_result = await self.import_protocol_services(file, protocol_name=None, operator=operator)
                if single_result and getattr(single_result, 'is_success', False):
                    success_items.append({'filename': filename, **(single_result.result or {})})
                else:
                    failed_items.append({'filename': filename, 'reason': getattr(single_result, 'message', '导入失败')})
            except Exception as exc:  # noqa: BLE001
                failed_items.append({'filename': filename, 'reason': str(exc)})

        status = ImportStatus.SUCCESS
        if failed_items and success_items:
            status = ImportStatus.PARTIAL
        elif failed_items and not success_items:
            status = ImportStatus.FAILED

        return ImportResult(
            status=status,
            total_count=len(files),
            success_count=len(success_items),
            failed_count=len(failed_items),
            data={'success': success_items, 'failed': failed_items},
        )

    @staticmethod
    def __read_workbook(content: bytes, suffix: str):
        try:
            if suffix in ('xlsx', 'xls'):
                return pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=str)
            text_io = io.StringIO(content.decode('utf-8-sig'))
            df = pd.read_csv(text_io, dtype=str)
            return {'字段列表': df}
        except Exception as exc:
            raise ServiceException(message=f'解析导入文件失败: {exc}') from exc

    @staticmethod
    def __parse_json_protocol(content: bytes) -> Dict[str, Any]:
        """
        解析 JSON 格式的协议配置文件

        JSON 格式示例:
        {
            "protocolType": "LAN",
            "protocolName": "协议名称",
            "header": {
                "sender": "发送方",
                "receiver": "接收方",
                "frequency": "1000",
                ...
            },
            "fields": [
                {
                    "code": "field1",
                    "name": "字段1",
                    "type": "string",
                    "seq": 1,
                    ...
                },
                ...
            ]
        }

        :param content: JSON 文件内容（字节）
        :return: 包含 header, fields, protocolType 的字典
        """
        try:
            # 尝试解码为 UTF-8
            text_content = content.decode('utf-8')
        except UnicodeDecodeError:
            try:
                # 尝试 UTF-8-sig (带 BOM)
                text_content = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise ServiceException(message='JSON 文件编码格式不支持，请使用 UTF-8 编码')

        try:
            data = json.loads(text_content)
        except json.JSONDecodeError as e:
            raise ServiceException(message=f'JSON 格式错误: {str(e)}')

        if not isinstance(data, dict):
            raise ServiceException(message='JSON 文件根节点必须是对象类型')

        # 提取协议类型
        protocol_type = data.get('protocolType') or data.get('protocol_type') or 'LAN'

        # 获取 protocolConfig 嵌套对象（如果存在）
        protocol_config = data.get('protocolConfig', {})
        if not isinstance(protocol_config, dict):
            protocol_config = {}

        # 提取头部配置
        header = {}
        # 支持多种可能的头部配置位置
        header_source = (
            data.get('header')
            or data.get('headerConfig')
            or data.get('messageConfig', {}).get('header')
            or protocol_config.get('header')
            or protocol_config.get('headerConfig')
            or
            # 如果 protocolConfig 本身包含配置字段（平铺结构），也尝试提取
            (protocol_config if protocol_config and 'fields' in protocol_config else None)
            or {}
        )

        if isinstance(header_source, dict):
            # 过滤掉 protocolType 和 fields 字段（如果存在）
            header = {k: v for k, v in header_source.items() if k not in ('protocolType', 'fields')}
        else:
            header = {}

        # 提取字段列表
        fields = []
        fields_source = (
            data.get('fields')
            or data.get('fieldsConfig')
            or data.get('messageConfig', {}).get('fields')
            or protocol_config.get('fields')
            or protocol_config.get('fieldsConfig')
            or []
        )

        if isinstance(fields_source, list):
            fields = fields_source
        elif isinstance(fields_source, dict):
            # 如果是字典，尝试转换为列表
            fields = list(fields_source.values()) if fields_source else []

        # 验证字段列表
        if not fields:
            raise ServiceException(message='JSON 文件中未找到有效的字段列表（fields）')

        # 确保字段列表中的每个元素都是字典
        validated_fields = []
        for idx, field in enumerate(fields):
            if isinstance(field, dict):
                validated_fields.append(field)
            else:
                raise ServiceException(message=f'字段列表第 {idx + 1} 项格式错误，必须是对象类型')

        return {'header': header, 'fields': validated_fields, 'protocolType': protocol_type}

    @staticmethod
    def __build_default_layout(header_fields: List[tuple]):
        layout: List[List[tuple]] = []
        buffer: List[tuple] = []
        for item in header_fields:
            buffer.append(item)
            if len(buffer) == 2:
                layout.append(buffer)
                buffer = []
        if buffer:
            layout.append(buffer)
        return layout

    @staticmethod
    def __header_sample_value(key: str):
        if key in DEFAULT_HEADER_SAMPLE_EXTRA:
            return DEFAULT_HEADER_SAMPLE_EXTRA[key]
        return DEFAULT_HEADER_SAMPLE.get(key, '')

    @staticmethod
    def __field_value(column_key: str, row_data: Dict, default_seq: int):
        if column_key == 'seq':
            return row_data.get('seq', default_seq)
        if column_key == 'parent_code':
            return row_data.get('parent_code') or row_data.get('parentId') or ''
        return row_data.get(column_key, '')

    @staticmethod
    def __normalize_header_value(key: str, value: Optional[str]):
        if key in HEADER_NUMERIC_KEYS:
            return ProtocolService.__safe_number(value)
        if value is None:
            return ''
        return str(value).strip()

    @staticmethod
    def __normalize_header_enums(header: Dict[str, Any], protocol_type: str) -> Dict[str, Any]:
        """
        将头部配置中的枚举/字典值从人类可读的标签（如 100Mbps、UDP、重传、BC-RT）
        归一化为系统字典的 value（如 100M、udp_unicast、timeout_retry、bc_to_rt）。
        """
        if not isinstance(header, dict):
            return header or {}

        def _normalize_key(text: str) -> str:
            return (
                str(text)
                .strip()
                .lower()
                .replace(' ', '')
                .replace('—', '-')
                .replace('－', '-')
                .replace('→', '->')
                .replace('至', '->')
            )

        def _map_value(raw_value: Any, mapping: Dict[str, str]) -> Any:
            if raw_value is None:
                return raw_value
            raw_text = str(raw_value).strip()
            if raw_text == '':
                return raw_value
            key = _normalize_key(raw_text)
            return mapping.get(key, raw_value)

        # 频率（跨协议共用）
        frequency_map = {
            'once': 'once',
            '单次': 'once',
            '10': '10',
            '10ms': '10',
            '10hz': '10',
            '20': '20',
            '20ms': '20',
            '20hz': '20',
            '50': '50',
            '50ms': '50',
            '50hz': '50',
            '100': '100',
            '100ms': '100',
            '100hz': '100',
            '160': '160',
            '160ms': '160',
            '200': '200',
            '200ms': '200',
            '500': '500',
            '500ms': '500',
            '1000': '1000',
            '1000ms': '1000',
            '1000hz': '1000',
        }

        # 速率（按协议区分）
        lan_speed_map = {
            '10m': '10M',
            '10mbps': '10M',
            '10mb/s': '10M',
            '100m': '100M',
            '100mbps': '100M',
            '100mb/s': '100M',
            '1g': '1000M',
            '1gbps': '1000M',
            '1000m': '1000M',
            '1000mbps': '1000M',
            '1000mb/s': '1000M',
        }
        b1553_speed_map = {
            '1m': '1M',
            '1mbps': '1M',
            '1mb/s': '1M',
        }

        # 传输方式（按协议区分）
        lan_method_map = {
            'udp': 'udp_unicast',
            'udp单播': 'udp_unicast',
            'udp_unicast': 'udp_unicast',
            'udp组播': 'udp_multicast',
            'udpmulticast': 'udp_multicast',
            'udp广播': 'udp_broadcast',
            'udpbroadcast': 'udp_broadcast',
            'tcp': 'tcp',
        }
        b1553_method_map = {
            'bc-rt': 'bc_to_rt',
            'bc->rt': 'bc_to_rt',
            'bc到rt': 'bc_to_rt',
            'rt-bc': 'rt_to_bc',
            'rt->bc': 'rt_to_bc',
            'rt到bc': 'rt_to_bc',
            'rt-rt': 'rt_to_rt',
            'rt->rt': 'rt_to_rt',
            'rt到rt': 'rt_to_rt',
        }

        # 错误处理（通用）
        error_handle_map = {
            'none': 'none',
            '不判断': 'none',
            'retry': 'timeout_retry',
            '重传': 'timeout_retry',
            '超时重传': 'timeout_retry',
            '超时重传3次,100ms': 'timeout_retry_3_100ms',
            '超时重传3次，100ms': 'timeout_retry_3_100ms',
            'crc': 'crc_check',
            'crc校验': 'crc_check',
            'parity_check': 'parity_check',
            '奇偶校验': 'parity_check',
            '-': 'default',
            'default': 'default',
        }

        proto = (protocol_type or '').upper()
        normalized = {**header}

        if 'frequency' in normalized:
            normalized['frequency'] = _map_value(normalized.get('frequency'), frequency_map)

        speed_map = lan_speed_map if proto == 'LAN' else b1553_speed_map if proto == '1553B' else None
        if speed_map and 'speed' in normalized:
            normalized['speed'] = _map_value(normalized.get('speed'), speed_map)
        # 兼容 baudRate 写入 speed
        if speed_map and 'baudRate' in normalized and not normalized.get('speed'):
            normalized['speed'] = _map_value(normalized.get('baudRate'), speed_map)

        if 'method' in normalized:
            method_map = lan_method_map if proto == 'LAN' else b1553_method_map if proto == '1553B' else None
            if method_map:
                normalized['method'] = _map_value(normalized.get('method'), method_map)

        # 错误处理：处理 errorHandling / errorHandle 两种写法
        if 'errorHandling' in normalized:
            normalized['errorHandling'] = _map_value(normalized.get('errorHandling'), error_handle_map)
        if 'errorHandle' in normalized:
            normalized['errorHandle'] = _map_value(normalized.get('errorHandle'), error_handle_map)
        # 同步两个 key，前端会优先使用 errorHandle
        if normalized.get('errorHandle') is None and normalized.get('errorHandling') is not None:
            normalized['errorHandle'] = normalized.get('errorHandling')
        if normalized.get('errorHandling') is None and normalized.get('errorHandle') is not None:
            normalized['errorHandling'] = normalized.get('errorHandle')

        return normalized

    @staticmethod
    def __iter_label_value_pairs(row: pd.Series, columns):
        """
        遍历行数据，尝试提取「标签-值」对，兼容字段1/值1、label/value 等双列布局
        """
        column_names = [str(col).strip() for col in columns]
        row_dict = {str(col).strip(): row.get(col) for col in columns}
        candidate_pairs = [
            ('字段1', '值1'),
            ('字段2', '值2'),
            ('label1', 'value1'),
            ('label2', 'value2'),
            ('字段', '值'),
            ('名称', '取值'),
        ]
        for label_col, value_col in candidate_pairs:
            if label_col in row_dict and value_col in row_dict:
                label = str(row_dict[label_col]).strip()
                if label != '':
                    yield label, row_dict[value_col]

        # 兼容第一列为字段名称、第二列为取值的简单两列表格（但跳过协议名称列）
        if len(column_names) >= 2:
            first_col = column_names[0]
            if first_col not in {'协议', 'Protocol', '协议名', '协议名称', 'protocolName', 'protocol'}:
                label = str(row_dict.get(first_col, '')).strip()
                if label:
                    yield label, row_dict.get(column_names[1], '')

    @staticmethod
    def __fill_header_from_matrix(df: pd.DataFrame, header_result: Dict[str, Optional[str]]):
        """
        扫描整个矩阵，兼容“标签/取值写在相邻列或上下单元格”的模板布局。
        """
        if df.empty:
            return

        matrix = df.fillna('').astype(str).values
        row_count, col_count = matrix.shape

        for row_index in range(row_count):
            for col_index in range(col_count):
                cell_value = str(matrix[row_index][col_index]).strip()
                if not cell_value:
                    continue
                key = HEADER_LABEL_TO_KEY.get(cell_value)
                if key is None:
                    continue

                current_value = header_result.get(key)
                if current_value not in (None, ''):
                    continue

                candidate = ''
                if col_index + 1 < col_count:
                    right_value = str(matrix[row_index][col_index + 1]).strip()
                    if right_value:
                        candidate = right_value

                if not candidate and row_index + 1 < row_count:
                    down_value = str(matrix[row_index + 1][col_index]).strip()
                    if down_value and down_value not in HEADER_LABEL_TO_KEY:
                        candidate = down_value

                if not candidate:
                    continue

                normalized_value = ProtocolService.__normalize_header_value(key, candidate)
                if normalized_value in (None, ''):
                    continue
                header_result[key] = normalized_value

    @staticmethod
    def __extract_protocol_type(sheet_map: Dict[str, pd.DataFrame]):
        """
        从报文配置sheet中提取协议类型

        同时兼容两种布局：
        1. 横向：第一行是表头，存在「协议类型」列，其首行数据为协议类型
        2. 纵向：第一列是字段名称，某一行的名称为「协议类型」，第二列为对应取值
        """
        config_df: Optional[pd.DataFrame] = None
        for sheet_name in HEADER_SHEET_NAMES:
            if sheet_name in sheet_map:
                config_df = sheet_map[sheet_name]
                break

        if config_df is None or config_df.empty:
            return None

        df = config_df.fillna('')
        col_aliases = {'协议类型', 'ProtocolType', 'Protocol Type'}
        display_reverse = {v.upper(): k.lower() for k, v in PROTOCOL_DISPLAY_NAME.items()}
        type_reverse_map = {v.upper(): k for k, v in PROTOCOL_TYPE_MAPPING.items()}

        # 方案一：横向表头，查找名为“协议类型”的列
        for column in df.columns:
            if str(column).strip() in col_aliases:
                series = df[column].astype(str).str.strip()
                for value in series:
                    if value:
                        protocol_type_raw = str(value).strip().upper()
                        return type_reverse_map.get(protocol_type_raw, protocol_type_raw.lower())

        # 方案二：纵向字段行，第一列作为字段名称，第二列存放值
        if len(df.columns) >= 2:
            name_col = df.columns[0]
            value_col = df.columns[1]
            for _, row in df.iterrows():
                field_name = str(row.get(name_col, '')).strip()
                if field_name in col_aliases:
                    value = str(row.get(value_col, '')).strip()
                    if value:
                        protocol_type_raw = value.upper()
                        return type_reverse_map.get(protocol_type_raw, protocol_type_raw.lower())

        # 方案三：字段/值双列布局（字段1/值1 等）
        for _, row in df.iterrows():
            for label, value in ProtocolService.__iter_label_value_pairs(row, df.columns):
                if str(label).strip() in col_aliases and value:
                    protocol_type_raw = str(value).strip().upper()
                    return type_reverse_map.get(protocol_type_raw, protocol_type_raw.lower())

        # 方案四：使用第一列的协议名称推断
        first_col_name = str(df.columns[0]).strip()
        if first_col_name in {'协议', 'Protocol', '协议名', '协议名称', 'protocolName', 'protocol'}:
            for _, row in df.iterrows():
                name_value = str(row.get(first_col_name, '')).strip()
                if not name_value:
                    continue
                mapped = display_reverse.get(name_value.upper())
                if mapped:
                    return mapped

        return None

    @staticmethod
    def __extract_header(sheet_map: Dict[str, pd.DataFrame]):
        header_df: Optional[pd.DataFrame] = None
        for sheet_name in HEADER_SHEET_NAMES:
            if sheet_name in sheet_map:
                header_df = sheet_map[sheet_name]
                break
        if header_df is None:
            first_sheet = next(iter(sheet_map.values()))
            if set(first_sheet.columns).issuperset({label for _, label in HEADER_FIELDS}):
                header_df = first_sheet
        header_result = {key: (None if key in HEADER_NUMERIC_KEYS else '') for key in ALL_HEADER_KEYS}
        if header_df is None or header_df.empty:
            return header_result

        df = header_df.fillna('')

        # 实际头部配置字段（排除元数据字段 '协议类型' 等）
        config_labels = {
            '发送方',
            '接收方',
            '传输频率',
            '传输速率/bps',
            '传输方式',
            '端口号',
            '发送时长/ms',
            '帧长度/Byte',
            '帧长度/word',
            '错误处理',
            '子地址',
            'ID',
        }

        # 方案一：横向表头（列名为中文标签）- 需要匹配到实际配置字段
        row_data = df.iloc[0].to_dict() if len(df) > 0 else {}
        matched_config_labels = [label for label in config_labels if label in row_data]
        if len(matched_config_labels) >= 3:  # 至少匹配到3个配置标签才认为是横向表头
            for label, key in HEADER_LABEL_TO_KEY.items():
                if label in row_data:
                    value = row_data.get(label, '')
                    header_result[key] = ProtocolService.__normalize_header_value(key, value)
            return header_result

        # 方案二：字段/值成对布局
        for _, row in df.iterrows():
            for label, value in ProtocolService.__iter_label_value_pairs(row, df.columns):
                key = HEADER_LABEL_TO_KEY.get(str(label).strip())
                if key is None:
                    continue
                header_result[key] = ProtocolService.__normalize_header_value(key, value)

        # 方案三：扫描矩阵，处理标签/取值在相邻单元格的布局（如导出的Excel格式）
        ProtocolService.__fill_header_from_matrix(df, header_result)
        return header_result

    @staticmethod
    def __extract_fields(sheet_map: Dict[str, pd.DataFrame]):
        fields_sheet: Optional[pd.DataFrame] = None
        for sheet_name in FIELD_SHEET_NAMES:
            if sheet_name in sheet_map:
                fields_sheet = sheet_map[sheet_name]
                break
        if fields_sheet is None:
            candidate = next(iter(sheet_map.values()))
            if any(header in candidate.columns for header in FIELD_COLUMN_ALIASES.get('field_name', [])):
                fields_sheet = candidate
        if fields_sheet is None:
            raise ServiceException(message='未找到包含字段定义的工作表，请确认模板未被修改')

        records = fields_sheet.fillna('').to_dict(orient='records')
        resolved_fields = []
        code_lookup: Dict[str, str] = {}
        name_lookup: Dict[str, str] = {}
        for index, row in enumerate(records):
            normalized = ProtocolService.__normalize_field_row(row, index, resolved_fields, code_lookup, name_lookup)
            if normalized:
                resolved_fields.append(normalized)

        return resolved_fields

    @staticmethod
    def __normalize_field_row(
        row: Dict,
        index: int,
        resolved_fields: List[Dict],
        code_lookup: Dict[str, str],
        name_lookup: Dict[str, str],
    ):
        normalized_row = {str(k).strip(): row.get(k) for k in row.keys()}
        field_name = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('field_name', []))
        field_code = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('field_code', []))
        if not field_name and not field_code:
            return None

        unique_id = str(field_code).strip() if field_code else ProtocolService.__generate_field_id(index)
        unique_id = str(unique_id).strip() or ProtocolService.__generate_field_id(index)
        byte_count = ProtocolService.__to_int(
            ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('byte_count', [])), default=1
        )
        byte_sequence = ProtocolService.__format_sequence(
            ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('byte_sequence', []))
        )
        value_range = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('value_range', [])) or ''
        unit = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('unit', [])) or ''
        data_type_raw = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('data_type', []))
        scale = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('scale', [])) or '1'
        remark = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('remark', [])) or ''
        parent_ref = ProtocolService.__pick_value(normalized_row, FIELD_COLUMN_ALIASES.get('parent_code', []))
        parent_id = ProtocolService.__resolve_parent_id(parent_ref, resolved_fields, code_lookup, name_lookup)

        if field_name:
            name_lookup[str(field_name).strip()] = unique_id
        if field_code:
            code_lookup[str(field_code).strip()] = unique_id

        return {
            'id': unique_id,
            'parentId': parent_id,
            'fieldName': str(field_name).strip() if field_name else '',
            'byteCount': byte_count,
            'byteSequence': str(byte_sequence).strip(),
            'valueRange': str(value_range).strip(),
            'unit': str(unit).strip(),
            'dataType': ProtocolService.__normalize_data_type(data_type_raw),
            'scale': str(scale).strip() or '1',
            'remark': str(remark).strip(),
        }

    @staticmethod
    def __pick_value(row: Dict, keys: List[str]):
        for key in keys:
            value = row.get(key)
            if value is None:
                continue
            text = str(value).strip()
            if text != '':
                return value
        return None

    @staticmethod
    def __generate_field_id(index: int):
        return f'field_{index}_{uuid.uuid4().hex}'

    @staticmethod
    def __normalize_data_type(value: Optional[str]):
        if not value:
            return 'STRING'
        value_str = str(value).strip()
        upper = value_str.upper()
        if upper in DATA_TYPE_OPTIONS:
            return upper
        return DATA_TYPE_ALIASES.get(value_str.lower(), 'STRING')

    @staticmethod
    def __safe_number(value: Optional[str]):
        if value is None:
            return None
        text = str(value).strip()
        if text == '':
            return None
        try:
            number = float(text)
            return int(number) if number.is_integer() else number
        except ValueError:
            return None

    @staticmethod
    def __to_int(value: Optional[str], default: int = 1):
        number = ProtocolService.__safe_number(value)
        return default if number is None else int(number)

    @staticmethod
    def __resolve_parent_id(
        parent_ref: Optional[str],
        resolved_fields: List[Dict],
        code_lookup: Dict[str, str],
        name_lookup: Dict[str, str],
    ):
        if parent_ref is None:
            return None
        ref_text = str(parent_ref).strip()
        if not ref_text:
            return None
        if ref_text in code_lookup:
            return code_lookup[ref_text]
        if ref_text in name_lookup:
            return name_lookup[ref_text]
        if ref_text.isdigit():
            ref_index = int(ref_text)
            if 0 <= ref_index < len(resolved_fields):
                return resolved_fields[ref_index]['id']
            if 1 <= ref_index <= len(resolved_fields):
                return resolved_fields[ref_index - 1]['id']
        return None

    @staticmethod
    def __format_sequence(value: Optional[str]):
        if value is None:
            return ''
        number = ProtocolService.__safe_number(value)
        if number is None:
            return str(value).strip()
        if isinstance(number, float) and number.is_integer():
            return str(int(number))
        return str(number)

    async def _validate_protocol(self, protocol_data: AddProtocolModel):
        protocol_type = (protocol_data.protocol_type or '').upper()
        if protocol_type == '1553B':
            if protocol_data.rt_address is not None and not self._validate_rt_address(protocol_data.rt_address):
                raise ValueError('RT地址必须在0-31之间')
            if protocol_data.word_count is not None and not self._validate_word_count(protocol_data.word_count):
                raise ValueError('字数必须在1-32之间')
        elif protocol_type == 'CAN':
            if protocol_data.baud_rate is not None and protocol_data.baud_rate not in {125000, 250000, 500000, 1000000}:
                raise ValueError('波特率无效')
            if protocol_data.can_id is not None and not self._validate_can_id(protocol_data.can_id, extended=False):
                raise ValueError('CAN ID超出范围')
        elif protocol_type == 'ETHERNET':
            if protocol_data.ip_address is not None and not self._validate_ip_address(protocol_data.ip_address):
                raise ValueError('IP地址格式不正确')
        elif protocol_type == 'RS422':
            return
        return

    @staticmethod
    def _validate_ip_address(ip_address: str) -> bool:
        try:
            import ipaddress

            ipaddress.ip_address(ip_address)
            return True
        except ValueError:
            return False

    @staticmethod
    def _validate_rt_address(rt_address: int) -> bool:
        return 0 <= rt_address <= 31

    @staticmethod
    def _validate_word_count(word_count: int) -> bool:
        return 1 <= word_count <= 32

    @staticmethod
    def _validate_can_id(can_id: int, extended: bool = False) -> bool:
        return 0 <= can_id <= (0x1FFFFFFF if extended else 0x7FF)

    def parse_1553b_message(self, message_data: Dict[str, Any]) -> Dict[str, Any]:
        command_word = message_data.get('command_word', 0)
        data_words = message_data.get('data_words', [])
        return {
            'rt_address': (command_word >> 11) & 0x1F,
            'tr_bit': (command_word >> 10) & 0x01,
            'sub_address': (command_word >> 5) & 0x1F,
            'word_count': command_word & 0x1F,
            'data_words': data_words,
            'status_word': message_data.get('status_word'),
        }

    def parse_can_frame(self, frame_data: Dict[str, Any]) -> Dict[str, Any]:
        return {
            'can_id': frame_data.get('can_id'),
            'dlc': frame_data.get('dlc'),
            'data': frame_data.get('data', []),
        }

    def parse_ethernet_packet(self, packet_data: Dict[str, Any]) -> Dict[str, Any]:
        payload = packet_data.get('payload', b'')
        payload_length = len(payload) if payload is not None else 0
        return {
            'src_ip': packet_data.get('src_ip'),
            'dst_ip': packet_data.get('dst_ip'),
            'src_port': packet_data.get('src_port'),
            'dst_port': packet_data.get('dst_port'),
            'protocol': packet_data.get('protocol'),
            'payload_length': payload_length,
        }
