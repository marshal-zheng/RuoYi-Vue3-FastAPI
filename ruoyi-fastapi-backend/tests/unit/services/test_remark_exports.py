"""Business export services derive literal Excel cells from authoritative documents."""

import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from common.remark_document_v1 import plain_to_remark
from modules.system.application.services.config_service import ConfigService
from modules.system.application.services.dict_data_service import DictDataService
from modules.system.application.services.dict_type_service import DictTypeService
from modules.system.application.services.post_service import PostService
from modules.system.application.services.role_service import RoleService
from modules.system.application.services.user_service import UserService

EXPORTERS = [
    UserService.export_user_list_services,
    RoleService.export_role_list_services,
    PostService.export_post_list_services,
    DictTypeService(None, None, None).export_dict_type_list_services,
    DictDataService(None, None).export_dict_data_list_services,
    ConfigService.export_config_list_services,
]


@pytest.mark.parametrize('exporter', EXPORTERS, ids=['user', 'role', 'post', 'dict-type', 'dict-data', 'config'])
async def test_remark_exports_derive_text_preserve_order_and_disable_formula_cells(exporter):
    root = Path(__file__).resolve().parents[4] / 'shared/remarks/remark-a-v1'
    nested = json.loads((root / 'valid/nested-list.json').read_text())
    records = [
        {
            'status': '0',
            'sex': '0',
            'isDefault': 'N',
            'configType': 'N',
            'dept': {'deptName': 'Dept'},
            'remark': plain_to_remark('=1+1\n\n 中文😀 ', identity_key='export:1:remark'),
            'remarkText': 'forged',
        },
        {
            'status': '0',
            'sex': '0',
            'isDefault': 'N',
            'configType': 'N',
            'dept': {'deptName': 'Dept'},
            'remark': nested,
            'remarkText': 'forged',
        },
    ]
    data = await exporter(records)
    worksheet = load_workbook(BytesIO(data), data_only=False).active
    index = next(cell.column for cell in worksheet[1] if cell.value == '备注')
    cells = [worksheet.cell(row=row, column=index) for row in [2, 3]]
    assert [cell.value for cell in cells] == ['=1+1\n\n 中文😀 ', '3. 第一\n  - 子项\n4. 第二']
    assert all(cell.data_type == 's' for cell in cells)
    # Projection must not replace the source document on the caller's records.
    assert isinstance(records[0]['remark'], dict)
    assert records[1]['remark'] == nested
